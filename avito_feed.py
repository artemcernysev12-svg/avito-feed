# -*- coding: utf-8 -*-
"""Ежедневный файл автозагрузки Авито «Мебель Фаворит» — наличие и сроки как на сайте.

Каждое утро (GitHub Actions, .github/workflows/avito-feed.yml):
  1. Берёт основу base/avito-base.xlsx (все активные объявления; обновляется с ПК скриптом avito_baza.py).
  2. Берёт с деплой-репо сайта свежие остатки (stock-data/stock.json — их ночью кладёт робот сайта из
     Яндекс.Почты), каталог (catalog.js) и составы комплектов (catalog-data/kitchen-comp.js, wardrobe-kits-mif.js).
  3. Каждому объявлению ставит «Доступность» и первые строки описания:
       корпус на складе            → «В наличии», ✅ «Есть на складе — привезём в Стерлитамак в пятницу, D месяц»;
       корпус не на складе         → «Под заказ», 🕐 «Под заказ — срок поставки 2–4 недели …»;
       комплект, все модули есть   → «В наличии», 📦 «Все модули комплекта есть на складе фабрики …»;
       комплект, часть модулей     → «Под заказ», 🧩 «Часть модулей есть на складе фабрики …»;
       диван                       → «Доступность» как в основе (салон подключим отдельно); у «Под заказ» — 🎨 текст.
     Пятница — как на сайте (site-ui.js favNextBatch): заявки до среды 16:00 по Стерлитамаку → эта пятница,
     после — следующая. Ложные фразы старых описаний («В НАЛИЧИИ НА СКЛАДЕ, привозим каждую неделю в Четверг»)
     вычищаются.
  4. Пишет avito.xlsx — его Авито забирает по ссылке.
ЗАЩИТА: остатки старше 3 дней или объявлений заметно меньше, чем во вчерашнем файле → файл НЕ обновляется
(Авито берёт вчерашний), код выхода 2.

python avito_feed.py --data <папка с данными> [--out avito.xlsx] [--today 2026-10-21T10:00] [--obrazcy 10]
python avito_feed.py --download …   — сначала скачать данные с GitHub (так делает робот)
"""
import argparse
import collections
import datetime
import json
import os
import re
import sys
import urllib.request

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)                                          # python -I не кладёт папку скрипта в путь
import kit_logic as K  # noqa: E402
RAW = "https://raw.githubusercontent.com/artemcernysev12-svg/mebel-favorit/main/"
FILES = ["catalog.js", "stock-data/stock.json", "catalog-data/kitchen-comp.js", "catalog-data/wardrobe-kits-mif.js",
         "catalog-data/wardrobe-coupe.js", "catalog-data/wardrobe-coupe-mif.js", "catalog-data/mattress-options.js"]
MATTRESS = {}       # матрасы: одна карточка — много размеров (app.js applyMattressStock)
COUPE = {}          # шкафы-купе Эры, комоды Клео… = корпус + фасад (app.js getWardrobeCoupeAvailability)
STERL = datetime.timezone(datetime.timedelta(hours=5))          # Стерлитамак, UTC+5
MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября",
          "ноября", "декабря"]
MAX_DESC = 7500                                                   # предел описания Авито (символов)

# ----- тексты (утверждены заказчиком 21.10.2026) -----
def T_STOCK(fri):
    return ("✅ Есть на складе — привезём в Стерлитамак в пятницу, " + fri + ".",
            "Доставка — в пятницу, субботу или другой удобный вам день, время согласуем заранее.",
            "Можно забрать самовывозом — где и когда, подскажем при заказе.")


T_ORDER = ("🕐 Под заказ — срок поставки 2–4 недели, точную дату назовём при оформлении.",
           "Самовывоз или доставка — день и время согласуем заранее.",
           "Нужно быстрее? Напишите — подберём похожие модели со склада или из салона.")


def T_KIT(fri):
    return ("📦 Все модули комплекта есть на складе фабрики — привезём в ближайшую пятницу, " + fri + ".",
            "Доставим в пятницу, субботу или другой удобный вам день по согласованию.",
            "Возможен самовывоз — подробности подскажем при заказе.")


def T_MATTRESS(sizes, fri, all_):
    tail = "" if all_ else " Остальные размеры — под заказ, 2–4 недели."
    return ("✅ Есть на складе: " + ", ".join(sizes) + " — привезём в Стерлитамак в пятницу, " + fri + "." + tail,
            "Доставка — в пятницу, субботу или другой удобный вам день, время согласуем заранее.",
            "Можно забрать самовывозом — где и когда, подскажем при заказе.")


T_KIT_PART = ("🧩 Часть модулей есть на складе фабрики, остальные привезём под заказ за 2–4 недели.",
              "Точную дату назовём при оформлении.",
              "Доставку или самовывоз согласуем заранее.")
T_SOFA_ORDER = ("🎨 Закажем диван в выбранном вами цвете — срок поставки 2–4 недели.",
                "Точную дату назовём при оформлении.",
                "Доставку или самовывоз согласуем заранее.")
OUR_EMOJI = ("✅ Есть на складе", "🕐 Под заказ", "📦 Все модули", "🧩 Часть модулей", "🎨 Закажем диван",
             "🛋 Диван есть в салоне")

# ----- ложные фразы старых описаний (наличие «на складе» и «привозим в четверг» писали всем подряд) -----
I = re.I
FIX = [
    (re.compile(r"\s*[—–-]?\s*В НАЛИЧИИ НА СКЛАДЕ\s*[,—–-]?\s*(?:привозим каждую неделю\s*(?:в|по)\s*четверг(?:ам)?"
                r"(?:\s*[-–]\s*пятницу)?\.?)?", I), ""),
    (re.compile(r"\s*[—–-]\s*наличие уточняйте в сообщении", I), ""),
    (re.compile(r"\s*Привозим каждую неделю\s*(?:<strong>)?\s*в\s*четверг\s*[–-]\s*пятницу\s*(?:</strong>)?\.?", I), ""),
    (re.compile(r"\s*привозим каждую неделю\s*(?:в|по)\s*четверг(?:ам)?(?:\s*[-–]\s*пятницу)?\.?", I), ""),
    (re.compile(r"\s+В НАЛИЧИИ(?=\s+ПО ДОСТУПНОЙ)", I), ""),
    (re.compile(r"(НОВЫЕ\s+[А-ЯЁA-Z]+)\s+В НАЛИЧИИ"), r"\1"),
    (re.compile(r"Кровати в наличии ограничены\.\s*(?:<br\s*/?>)?\s*", I), ""),
    (re.compile(r"Поставк[аи] 1 раз в неделю\s*[—–-]\s*каждый четверг\.\s*(?:<br\s*/?>)?\s*", I), ""),
    (re.compile(r"\s*\(четверг\)", I), ""),
    (re.compile(r"Также в наличии другие", I), "Также есть другие"),
    (re.compile(r"\s*В наличии\.\s*Ограниченное количество\.", I), ""),
]
CLEAN = [
    (re.compile(r"(?:<br\s*/?>\s*)+(</strong>)"), r"\1"),
    (re.compile(r"<strong>\s*</strong>"), ""),
    (re.compile(r"(<p>)\s*(?:<br\s*/?>\s*)+"), r"\1"),
    (re.compile(r"(?:<br\s*/?>\s*)+(</p>)"), r"\1"),
    (re.compile(r"(?:<br\s*/?>\s*){3,}"), "<br> <br> "),
    (re.compile(r"<p>\s*</p>\s*"), ""),
]


def ru_date(d):
    return f"{d.day} {MONTHS[d.month - 1]}"


def next_friday(now):
    """site-ui.js favNextBatch: ближайшая среда 16:00 (если прошла — следующая) + 2 дня."""
    wed = (now + datetime.timedelta(days=(2 - now.weekday()) % 7)).replace(hour=16, minute=0, second=0, microsecond=0)
    if now >= wed:
        wed += datetime.timedelta(days=7)
    return (wed + datetime.timedelta(days=2)).date()


def js_obj_after(path, marker, arr=False):
    s = open(path, encoding="utf-8").read()
    i = s.find(marker)
    if i < 0:
        return {}
    j = s.index("[" if arr else "{", i + len(marker))
    return json.JSONDecoder().raw_decode(s[j:])[0]


def download(dst):
    for f in FILES:
        p = os.path.join(dst, *f.split("/"))
        os.makedirs(os.path.dirname(p), exist_ok=True)
        last = None
        for _ in range(3):
            try:
                with urllib.request.urlopen(RAW + f, timeout=120) as r:
                    open(p, "wb").write(r.read())
                break
            except Exception as e:                       # noqa: BLE001 — сеть: пробуем ещё
                last = e
        else:
            sys.exit(f"не скачался {f}: {last}")


def load_data(d):
    cat = js_obj_after(os.path.join(d, "catalog.js"), "window.CATALOG=", arr=True)
    K.CAT_BY_ID.clear()
    K.CAT_BY_ID.update({str(x.get("id")): x for x in cat})
    comp = js_obj_after(os.path.join(d, "catalog-data", "kitchen-comp.js"), "__KITCHEN_COMP__ =")
    comp.update(js_obj_after(os.path.join(d, "catalog-data", "wardrobe-kits-mif.js"),
                             "Object.assign(window.__KITCHEN_COMP__ || {},"))
    K.KIT_COMP.clear()
    K.KIT_COMP.update(comp)
    MATTRESS.clear()
    MATTRESS.update(js_obj_after(os.path.join(d, "catalog-data", "mattress-options.js"), "__MATTRESS_OPTIONS__"))
    COUPE.clear()
    COUPE.update(js_obj_after(os.path.join(d, "catalog-data", "wardrobe-coupe.js"), "window.__WARDROBE_COUPE__"))
    COUPE.update(js_obj_after(os.path.join(d, "catalog-data", "wardrobe-coupe-mif.js"),
                              "Object.assign(window.__WARDROBE_COUPE__ || {},"))
    st = json.load(open(os.path.join(d, "stock-data", "stock.json"), encoding="utf-8"))
    stock = {K.norm_art(k): v for k, v in (st.get("map") or {}).items()}
    dd, mm, yy = map(int, st["date"].split("."))
    return stock, datetime.date(yy, mm, dd)


def coupe_qty(rec, stock):
    """app.js getWardrobeCoupeAvailability: минимум по компонентам (stock «Нет» у компонента — не считаем)."""
    need, missing = {}, False
    for c in rec.get("components") or []:
        if str(c.get("stock") or "").lower() == "нет":
            continue
        art = str(c.get("art") or "").strip() or str((K.CAT_BY_ID.get(str(c.get("id") or "")) or {}).get("art") or "")
        if not art:
            missing = True
            continue
        k = K.norm_art(art)
        need[k] = need.get(k, 0) + max(1, int(K._kit_num(c.get("qty")) or 1))
    if missing or not need:
        return 0
    q = []
    for k, nd in need.items():
        if k not in stock:
            return 0
        q.append(int(float(stock[k] or 0)) // nd)
    return max(0, min(q))


def mattress_sizes(rec, stock):
    """app.js mattressVariantQty: размер есть, если по любому его артикулу (v.arts или v.art) остаток > 0."""
    have, alls = [], []
    for v in rec.get("variants") or []:
        sz = str(v.get("size") or "").replace("x", "×")
        if sz not in alls:
            alls.append(sz)
        arts = [x.get("a") for x in (v.get("arts") or [])] or [v.get("art")]
        if sum(K.stock_by_art(a, stock) for a in arts if a) > 0 and sz not in have:
            have.append(sz)
    return have, alls


def decide(n, cur_av, M, stock, fri=""):
    """→ (доступность, тексты, вид для отчёта)."""
    info = M.get(n) or {}
    it = K.CAT_BY_ID.get(n)
    kind = info.get("kind") or "item"
    if kind == "sofa" or (it and it.get("c") == "Диваны"):
        if cur_av == "В наличии":
            return cur_av, None, "диван: как было (салон — следующим шагом)"
        return "Под заказ", T_SOFA_ORDER, "диван под заказ"
    if MATTRESS.get(n):
        have, alls = mattress_sizes(MATTRESS[n], stock)
        if have:
            return "В наличии", T_MATTRESS(have, fri, len(have) == len(alls)), f"матрас: на складе {len(have)} размер(ов)"
        return "Под заказ", T_ORDER, "матрас: под заказ"
    if COUPE.get(n):                                  # сайт проверяет купе/комоды «корпус+фасад» первыми
        q = coupe_qty(COUPE[n], stock)
        if q > 0:
            return "В наличии", "STOCK", f"корпус+фасад на складе ({q} шт.)"
        return "Под заказ", T_ORDER, "корпус+фасад: под заказ"
    if it and K.KIT_COMP.get(n):
        K.kit_stock.last = (0, 0)
        have, kits = K.kit_stock(it, stock)
        units, total = K.kit_stock.last
        if have:
            return "В наличии", "KIT", f"комплект: все модули есть ({kits} шт.)"
        if units:
            return "Под заказ", T_KIT_PART, f"комплект: часть модулей ({units} из {total})"
        return "Под заказ", T_ORDER, "комплект: под заказ"
    q = K.qty_of(it, stock) if it else K.stock_by_art(info.get("art"), stock)
    if q > 0:
        return "В наличии", "STOCK", f"на складе ({q} шт.)"
    if kind == "kit":
        return "Под заказ", T_ORDER, "комплект без состава на сайте: под заказ"
    return "Под заказ", T_ORDER, "под заказ" + ("" if (it or {}).get("art") or info.get("art") else " (нет артикула)")


def new_desc(old, lines):
    t = str(old or "")
    # наш блок уже стоит (вчерашний файл / повторный прогон) — снимаем, чтобы не двоился
    t = re.sub(r"^\s*<p><strong>(?:" + "|".join(map(re.escape, OUR_EMOJI)) + r").*?</p>\s*", "", t, flags=re.S)
    for rx, rep in FIX + CLEAN:
        t = rx.sub(rep, t)
    if not lines:
        return t.strip()
    block = f"<p><strong>{lines[0]}</strong><br> {lines[1]}<br> {lines[2]}</p> "
    return (block + t.strip()).strip()


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default=os.path.join(HERE, "_data"))
    ap.add_argument("--download", action="store_true")
    ap.add_argument("--base", default=os.path.join(HERE, "base", "avito-base.xlsx"))
    ap.add_argument("--map", default=os.path.join(HERE, "base", "avito-map.json"))
    ap.add_argument("--out", default=os.path.join(HERE, "avito.xlsx"))
    ap.add_argument("--today", help="для проверки: 2026-10-22T17:00 (время Стерлитамака)")
    ap.add_argument("--obrazcy", type=int, default=0, help="показать N примеров каждого вида")
    ap.add_argument("--max-age", type=int, default=3, help="остатки старше N дней — файл не обновляем")
    a = ap.parse_args()
    if a.download:
        download(a.data)
    now = datetime.datetime.fromisoformat(a.today) if a.today else datetime.datetime.now(STERL).replace(tzinfo=None)
    fri = ru_date(next_friday(now))
    stock, sdate = load_data(a.data)
    age = (now.date() - sdate).days
    print(f"сегодня {now:%d.%m.%Y %H:%M} (Стерлитамак) → ближайший привоз: пятница, {fri}; остатки от {sdate:%d.%m.%Y}"
          f" ({len(stock)} артикулов)")
    if age > a.max_age:
        print(f"!! остатки устарели на {age} дн. — файл НЕ обновляю")
        sys.exit(2)
    M = json.load(open(a.map, encoding="utf-8"))["map"]
    wb = openpyxl.load_workbook(a.base)
    rep, av_cnt, samples, long_ = collections.Counter(), collections.Counter(), collections.defaultdict(list), []
    total = 0
    for ws in wb.worksheets:
        if ws.title.startswith("Спр") or ws.title == "Инструкция":
            continue
        h = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        if "Номер объявления на Авито" not in h:
            continue
        cn, cd, ca = (h.index(x) + 1 for x in ("Номер объявления на Авито", "Описание объявления", "Доступность"))
        ct = h.index("Название объявления") + 1
        for r in range(5, ws.max_row + 1):
            n = str(ws.cell(r, cn).value or "").strip()
            if not n:
                continue
            total += 1
            old = ws.cell(r, cd).value
            av, lines, why = decide(n, ws.cell(r, ca).value, M, stock, fri)
            lines = T_STOCK(fri) if lines == "STOCK" else T_KIT(fri) if lines == "KIT" else lines
            d = new_desc(old, lines)
            if len(d) > MAX_DESC:                           # не влезает — без нашего блока, только чистка
                long_.append(n)
                d = new_desc(old, None)
            ws.cell(r, ca).value = av
            ws.cell(r, cd).value = d
            key = why.split(" (")[0]
            rep[key] += 1
            av_cnt[av] += 1
            if len(samples[key]) < a.obrazcy:
                samples[key].append((n, ws.cell(r, ct).value, why, old, d))
    print(f"объявлений: {total}; доступность: {dict(av_cnt)}")
    for k, v in rep.most_common():
        print(f"   {v:5}  {k}")
    if long_:
        print(f"   длинные описания (блок не влез, только чистка): {len(long_)} — {long_[:10]}")
    prev = 0
    if os.path.exists(a.out):
        wbp = openpyxl.load_workbook(a.out, read_only=True)
        for ws in wbp.worksheets:
            if not ws.title.startswith("Спр") and ws.title != "Инструкция":
                ws.reset_dimensions()
                prev += sum(1 for x in ws.iter_rows(min_row=5, max_col=1, values_only=True) if x and x[0])
        wbp.close()                                   # иначе Windows не даст заменить файл
    if prev and total < prev * 0.95:
        print(f"!! объявлений {total}, а вчера было {prev} — файл НЕ обновляю (проверить основу)")
        sys.exit(2)
    for k, lst in samples.items():
        for n, t, why, old, d in lst:
            print(f"\n=== {k} | {n} | {t} | {why}\n--- было:  {re.sub(r'<[^>]+>', ' ', str(old))[:300]}"
                  f"\n--- стало: {re.sub(r'<[^>]+>', ' ', d)[:420]}")
    tmp = a.out + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, a.out)
    print("\nзаписано:", a.out)


if __name__ == "__main__":
    main()
